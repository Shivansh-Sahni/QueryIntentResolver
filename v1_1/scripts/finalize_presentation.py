"""Publish the explicit no-promotion decision without changing model or test outputs."""
from __future__ import annotations
import json
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

ROOT=Path(__file__).resolve().parents[1]
A=ROOT/'artifacts'; D=ROOT/'docs'
s=json.loads((A/'SUMMARY.json').read_text(encoding='utf-8'))
b=s['benchmark']; legacy=s['legacy']['v1_0_1']; app=s['applicability']['core']; review=s['review']

def pct(value): return 'undefined' if value is None else f'{100*value:.2f}%'

# Acceptance of executed evidence is separate from training-only model selection.
decision={
    'engineering_handoff_complete':True,
    'model_quality_objective_complete':bool(s['model_quality_passes_declared_gates']),
    'automatic_promotion_allowed':False,
    'production_authorized':False,
    'retained_reference':'v1.0.1 (safer historical point estimate; still provisional)',
    'v1_1_role':'research and integration candidate; not a replacement for the reference',
    'reason':'No live promotion is authorized. The executed V1.1 candidate fails the unchanged short-circuit safety requirement; aggregate accuracy does not waive it.',
    'historical_v1_1_false_sc_fraction':f"{b['false_short_circuit_count']}/{b['predicted_short_circuit_count']}",
    'historical_reference_false_sc_fraction':f"{legacy['false_short_circuit_count']}/{legacy['predicted_short_circuit_count']}",
    'paired_accuracy_gain_95ci':s['paired_accuracy']['v1_0_1']['bootstrap_95ci'],
    'small_sample_warning':'Neither point estimate certifies live safety; the old reference also misses its short-circuit recall target.',
    'remaining_gates':['short-circuit safety without test tuning','independent real-query validation','approved handler bindings and index coverage','product-owner deployment approval'],
}
(A/'RELEASE_DECISION.json').write_text(json.dumps(decision,indent=2)+'\n',encoding='utf-8')
base=json.loads((A/'COMPLETION_MATRIX.json').read_text())
base.update(automatic_promotion_allowed=False,production_authorized=False,
            model_quality_objective_complete=bool(s['model_quality_passes_declared_gates']),
            promotion_decision='retain_v1_0_1_as_safer_provisional_reference; V1.1 is not promoted')
(A/'COMPLETION_MATRIX.json').write_text(json.dumps(base,indent=2,sort_keys=True)+'\n')
statement=(f"The standalone engineering handoff is complete. The production-quality objective is not. "
           f"Do not automatically promote V1.1: it falsely short-circuited {b['false_short_circuit_count']} of "
           f"{b['predicted_short_circuit_count']} predicted short circuits ({pct(b['false_short_circuit_rate'])}), "
           f"versus {legacy['false_short_circuit_count']} of {legacy['predicted_short_circuit_count']} "
           f"({pct(legacy['false_short_circuit_rate'])}) for the preserved V1.0.1 reference. "
           "Keep the original reference for comparison and V1.1 for research/integration review. Neither is authorized for live traffic.")
delta=s['paired_accuracy']['v1_0_1']
body=f'''# Release Decision: Engineering Delivered, No Automatic Promotion

{statement}

## Evidence, not a deadline-driven release

The observed V1.1 accuracy gain is {100*delta['candidate_minus_reference']:.2f} percentage points on the old 300-query regression benchmark, but its paired row-bootstrap 95% interval is {100*delta['bootstrap_95ci'][0]:.2f} to {100*delta['bootstrap_95ci'][1]:.2f} points. Training conditions differ: V1.1 excludes heuristic template families present in that benchmark. This does not establish a general improvement. The balanced assistant-authored core applicability suite achieved {pct(app['accuracy'])}; it is not independent real-user validation.

## Completed internal deliverables

- All 55 received review rows matched original quarantine identities. {review['accepted_unchanged']} labels were retained, {review['policy_adjudicated']} separately AI-policy-adjudicated, and {review['deferred_context']} context-dependent rows deliberately excluded. Original reviewer values remain intact.
- Three fixed candidates were evaluated using training-only grouped OOF selection and nested grouped calibration. The selected policy was saved before evaluation labels were read.
- Applicability execution, error reporting, 42 unit tests, 134 runtime checks, actual HTTP smoke testing and a Docker build/inference test are complete.
- Reproducible source, candidate model, hashes, full predictions, review audit, documentation and presentation files are delivered.

## Not represented as complete

The candidate does not meet the unchanged short-circuit safety target. Independent real traffic, actual MascotGO route bindings, supported index fields, acceptance tolerances and deployment authorization remain unfulfilled. Issue #5 is the external architecture decision record; the model-quality risk is separately retained here rather than falsely relabeled as an external-only dependency.

## Presentation order

1. This release decision.
2. REVIEWER_BRIEF.html and EXECUTIVE_HANDOFF.pdf.
3. TECHNICAL_CLOSEOUT.md, model comparison and applicability failures.
4. Local CLI/API demonstration.
5. INTEGRATION.md: concrete MascotGO decisions.

All v1/ files are preserved. No model, threshold, parser or test label changed in this presentation pass.
'''
(D/'RELEASE_DECISION.md').write_text(body,encoding='utf-8')
p=D/'TECHNICAL_CLOSEOUT.md'; t=p.read_text(encoding='utf-8')
marker='## Explicit promotion decision\n'
if marker not in t:
    pos=t.index('## Decision for the handoff')
    p.write_text(t[:pos]+marker+'\n'+statement+'\n\n'+t[pos:],encoding='utf-8')
p=D/'REVIEWER_BRIEF.html'; t=p.read_text(encoding='utf-8')
marker='<section id="promotion-decision">'
if marker not in t:
    pos=t.index('<section>')
    panel=marker+'<h2>Release decision: no automatic promotion</h2><p>'+escape(statement)+'</p><p><a href="RELEASE_DECISION.md">Decision and remaining quality gates</a></p></section>'
    p.write_text(t[:pos]+panel+t[pos:],encoding='utf-8')
# Word layout refinement without data or label changes.
p=D/'EXECUTIVE_HANDOFF.docx'; doc=Document(p)
for para in doc.paragraphs:
    para.paragraph_format.keep_together=True
    if para.text=='Applicability and runtime': para.paragraph_format.page_break_before=True
    if para.text.startswith('The query-only engineering package is complete.'):
        para.text=statement
        para.paragraph_format.keep_together=True
footer=doc.sections[0].footer.paragraphs[0]
footer.text='Query Intent Resolver | Engineering handoff, not production authorization'
for run in footer.runs: run.font.size=Pt(8)
doc.core_properties.title='Query Intent Resolver - V1.1 technical closeout'
doc.core_properties.subject='Engineering evidence and explicit no-promotion decision'
doc.core_properties.author='Shivansh Sahni - AI-assisted closeout'
doc.save(p)
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='DecisionBody',fontName='Helvetica',fontSize=10,leading=14,spaceAfter=9))
styles.add(ParagraphStyle(name='DecisionHeading',fontName='Helvetica-Bold',fontSize=12,leading=16,spaceBefore=9,spaceAfter=6))
items=[Paragraph('Query Intent Resolver',styles['Title']),Paragraph('Engineering closeout | Release decision',styles['Normal']),Spacer(1,16),Paragraph('No automatic model promotion',styles['DecisionHeading']),Paragraph(escape(statement),styles['DecisionBody'])]
rows=[['System / set','Accuracy','False SC / predicted SC','SC recall'],
      ['V1.0.1 historical reference',pct(legacy['accuracy']),f"{legacy['false_short_circuit_count']}/{legacy['predicted_short_circuit_count']} ({pct(legacy['false_short_circuit_rate'])})",pct(legacy['short_circuit_recall'])],
      ['V1.1 historical regression',pct(b['accuracy']),f"{b['false_short_circuit_count']}/{b['predicted_short_circuit_count']} ({pct(b['false_short_circuit_rate'])})",pct(b['short_circuit_recall'])],
      ['V1.1 core applicability',pct(app['accuracy']),pct(app['false_short_circuit_rate']),pct(app['short_circuit_recall'])]]
table=Table(rows,colWidths=[169,75,147,81],repeatRows=1)
table.setStyle(TableStyle([('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),8),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#edf1f5')),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9),('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor('#d6dce3'))]))
items+=[table,Spacer(1,12),Paragraph('Completed and reproducible',styles['DecisionHeading']),Paragraph('55 review submissions processed with unchanged originals, separately attributed adjudications and context exclusions. Three fixed model families compared using grouped training-only selection. 120 diagnostic cases executed. 42 unit tests, 134 runtime checks, real HTTP and Docker checks passed. All original V1 files and the 300-query benchmark are preserved.',styles['DecisionBody']),Paragraph('How to interpret the evidence',styles['DecisionHeading']),Paragraph('The old benchmark has been observed before and is regression evidence, not a new blind test. Training conditions differ. The 100-case core applicability set is assistant-authored; 20 boundary cases are reported separately. The paired accuracy-gain interval includes zero. Neither a successful build nor a favorable point estimate proves production safety.',styles['DecisionBody']),Paragraph('What remains open',styles['DecisionHeading']),Paragraph('Short-circuit safety improvement remains a model-quality requirement. Peter/MascotGO must confirm actual route handlers, index coverage, invocation surfaces, deployment environment, authorized real-query validation and risk tolerances. No live service or Foundry integration is claimed.',styles['DecisionBody']),Paragraph('Presentation materials',styles['DecisionHeading']),Paragraph('Repository: Shivansh-Sahni/QueryIntentResolver. Start with v1_1/docs/REVIEWER_BRIEF.html and TECHNICAL_CLOSEOUT.md. Exact evidence is in v1_1/artifacts/SUMMARY.json, RELEASE_DECISION.json and FINAL_VALIDATION.json.',styles['DecisionBody'])]
def foot(canvas,doc):
    canvas.setFont('Helvetica',8)
    canvas.drawString(54,30,'Engineering delivery complete | Model quality and live deployment remain gated')
    canvas.drawRightString(558,30,str(doc.page))
SimpleDocTemplate(str(D/'RELEASE_DECISION.pdf'),pagesize=letter,leftMargin=54,rightMargin=54,topMargin=46,bottomMargin=46).build(items,onFirstPage=foot,onLaterPages=foot)
print('Presentation finalized without model/test changes')
